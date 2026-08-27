import re
import os
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


BASE_DIR = Path(__file__).resolve().parent
MD_PATH = BASE_DIR / "Informe-Final-EF.md"
OUT_PATH = Path(os.environ.get("INFORME_DOCX_OUT", BASE_DIR / "Informe-Final-EF.docx"))

H1_COLOR = RGBColor(0x1F, 0x38, 0x64)   # igual al T3
H2_COLOR = RGBColor(0x2C, 0x3E, 0x50)
H3_COLOR = RGBColor(0x34, 0x49, 0x5E)
GRAY = RGBColor(85, 85, 85)
LIGHT_GRAY = "F2F4F7"
BORDER = "B8C2CC"
BASE_FONT = "Aptos"
COVER_FONT = "Aptos Display"
MIN_FONT_SIZE = 11


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, name=BASE_FONT, size=None, color=None, bold=None, italic=None):
    if size is None or size < MIN_FONT_SIZE:
        size = MIN_FONT_SIZE
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size=MIN_FONT_SIZE, color=None, bold=None):
    if size is None or size < MIN_FONT_SIZE:
        size = MIN_FONT_SIZE
    style.font.name = BASE_FONT
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), BASE_FONT)
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def enforce_docx_typography(path):
    tmp_path = path.with_suffix(".tmp.docx")
    with ZipFile(path, "r") as source, ZipFile(tmp_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                try:
                    root = ET.fromstring(data)
                except ET.ParseError:
                    target.writestr(item, data)
                    continue
                changed = False
                for r_fonts in root.iter(qn("w:rFonts")):
                    # Preservar Aptos Display en portada; sólo normalizar el resto
                    current = r_fonts.get(qn("w:ascii"), "")
                    if current == COVER_FONT:
                        continue
                    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                        r_fonts.set(qn(f"w:{attr}"), BASE_FONT)
                    changed = True
                for tag in ("w:sz", "w:szCs"):
                    for elem in root.iter(qn(tag)):
                        value = elem.get(qn("w:val"))
                        if value and int(value) < MIN_FONT_SIZE * 2:
                            elem.set(qn("w:val"), str(MIN_FONT_SIZE * 2))
                            changed = True
                if changed:
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    tmp_path.replace(path)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    for style in styles:
        try:
            current_size = style.font.size.pt if style.font.size is not None else MIN_FONT_SIZE
            set_style_font(style, size=max(current_size, MIN_FONT_SIZE))
        except Exception:
            continue

    normal.font.name = BASE_FONT
    set_style_font(normal, size=MIN_FONT_SIZE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    BLACK = RGBColor(0, 0, 0)
    for style_name, size, before, after in [
        ("Heading 1", 16, 18, 5),
        ("Heading 2", 13, 12, 4),
        ("Heading 3", 12,  8, 3),
    ]:
        style = styles[style_name]
        set_style_font(style, size=size, color=BLACK, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    # Sin pie de página
    section.footer.is_linked_to_previous = False


def _cover_para(doc, text, size, bold=False, space_before=0, space_after=8):
    """Párrafo centrado con Aptos Display para la portada."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = COVER_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    from docx.oxml.ns import qn
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{attr}"), COVER_FONT)
    return p


def add_cover(doc, lines):
    # Extraer metadatos del markdown
    meta = {}
    for line in lines:
        m = re.match(r"- \*\*(.+?):\*\* (.*)", line)
        if m:
            meta[m.group(1).strip()] = m.group(2).strip()

    # Institución
    _cover_para(doc, "Universidad Privada del Norte (UPN)", size=14, bold=True, space_before=60, space_after=4)
    facultad = meta.get("Facultad", "Facultad de Ingeniería")
    carrera  = meta.get("Carrera", "Ingeniería de Sistemas Computacionales")
    _cover_para(doc, f"{facultad} — {carrera}", size=11, space_after=16)

    # Título del documento y del proyecto
    _cover_para(doc, "Informe Final del Proyecto – Examen Final", size=10, space_after=6)
    titulo = meta.get("Título del proyecto", "Sistema de gestión de clientes y ventas")
    _cover_para(doc, titulo, size=20, bold=True, space_after=20)

    # Campos: label en negrita + valor en normal, centrados
    campos = [
        ("Curso",         meta.get("Curso", "")),
        ("Organización",  meta.get("Organización seleccionada", "")),
        ("Integrante",    meta.get("Integrante", "")),
    ]
    for label, value in campos:
        if not value:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        r_label = p.add_run(f"{label}:")
        r_label.font.name = COVER_FONT
        r_label.font.bold = True
        r_label.font.size = Pt(11)
        rpr = r_label._element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn(f"w:{attr}"), COVER_FONT)
        pv = doc.add_paragraph()
        pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(4)
        r_val = pv.add_run(value)
        r_val.font.name = COVER_FONT
        r_val.font.size = Pt(11)
        rpr2 = r_val._element.get_or_add_rPr()
        rf2 = rpr2.get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf2.set(qn(f"w:{attr}"), COVER_FONT)

    # Ciclo al final
    _cover_para(doc, meta.get("Ciclo académico", "2026-1"), size=11, space_before=16, space_after=0)

    doc.add_page_break()


def parse_tables(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if re.match(r"^\|[\s:\-]+\|", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def add_markdown_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    for idx, row_data in enumerate(rows):
        for col in range(cols):
            cell = table.rows[idx].cells[col]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            text = row_data[col] if col < len(row_data) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=MIN_FONT_SIZE, bold=(idx == 0))
    doc.add_paragraph()


def resolve_image(path_text):
    path = Path(path_text)
    if path.is_absolute():
        return path
    return (BASE_DIR / path).resolve()


def add_image(doc, alt, path_text):
    img_path = resolve_image(path_text)
    if not img_path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[Imagen no encontrada: {path_text}]")
        set_run_font(run, color=RGBColor(155, 28, 28), italic=True)
        return

    try:
        with Image.open(img_path) as image:
            width_px, height_px = image.size
    except Exception:
        width_px, height_px = 1600, 900

    aspect = width_px / max(height_px, 1)
    width_inches = 6.25
    if aspect > 2.2:
        width_inches = 6.45
    elif aspect < 1.0:
        width_inches = 4.7

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(width_inches))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    r = caption.add_run(alt)
    set_run_font(r, size=MIN_FONT_SIZE, color=GRAY, italic=True)


def add_code_block(doc, code_lines):
    for code_line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_line.rstrip())
        set_run_font(run, name="Courier New", size=MIN_FONT_SIZE)
    doc.add_paragraph()


def clean_inline(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("---", "")
    return text.strip()


def add_paragraph_from_markdown(doc, line):
    stripped = line.strip()
    if not stripped:
        return
    if stripped == "---":
        return

    bullet = re.match(r"^- (.+)", stripped)
    numbered = re.match(r"^(\d+)\. (.+)", stripped)

    if bullet:
        p = doc.add_paragraph(style="List Bullet")
        text = bullet.group(1)
    elif numbered:
        # Usar Normal + número explícito para evitar numeración continua de Word
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        text = numbered.group(1) + ". " + numbered.group(2)
    else:
        p = doc.add_paragraph()
        text = stripped

    text = clean_inline(text)
    if not text:
        return

    run = p.add_run(text)
    set_run_font(run)


def build_document():
    md_lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    # Cover page from the Portada block
    cover_lines = []
    i = 0
    while i < len(md_lines):
        if md_lines[i].strip() == "## Portada":
            i += 1
            while i < len(md_lines) and md_lines[i].strip() != "---":
                cover_lines.append(md_lines[i])
                i += 1
            break
        i += 1
    add_cover(doc, cover_lines)

    # Continue after the cover block separator
    while i < len(md_lines) and md_lines[i].strip() != "---":
        i += 1
    if i < len(md_lines):
        i += 1

    code_block = None
    while i < len(md_lines):
        line = md_lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if code_block is None:
                code_block = []
            else:
                add_code_block(doc, code_block)
                code_block = None
            i += 1
            continue

        if code_block is not None:
            code_block.append(line)
            i += 1
            continue

        img_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if img_match:
            add_image(doc, img_match.group(1), img_match.group(2))
            i += 1
            continue

        if stripped.startswith("|"):
            rows, new_i = parse_tables(md_lines, i)
            add_markdown_table(doc, rows)
            i = new_i
            continue


        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 1"]
            p.add_run(clean_inline(stripped[2:]))
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 1"]
            p.add_run(clean_inline(stripped[3:]))
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 2"]
            p.add_run(clean_inline(stripped[4:]))
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.style = doc.styles["Heading 3"]
            p.add_run(clean_inline(stripped[5:]))
        else:
            add_paragraph_from_markdown(doc, line)
        i += 1

    doc.save(OUT_PATH)
    enforce_docx_typography(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
